import io
import datetime
import calendar
import pandas as pd
import streamlit as st
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==========================================
# KONFIGURASI HALAMAN STREAMLIT
# ==========================================
st.set_page_config(
    page_title="Kalkulator Gaji & Lembur (PP 35/2021)",
    page_icon="💼",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom Styling CSS untuk Antarmuka Modern & Premium
st.markdown("""
<style>
    /* Global Styling */
    .main {
        background-color: #f8fafc;
    }
    
    /* Hero Header */
    .hero-container {
        background: linear-gradient(135deg, #1e3a8a 0%, #0f172a 100%);
        color: white;
        padding: 1.8rem 2.2rem;
        border-radius: 16px;
        margin-bottom: 1.8rem;
        box-shadow: 0 10px 25px -5px rgba(15, 23, 42, 0.15);
    }
    .hero-title {
        font-size: 1.85rem;
        font-weight: 800;
        margin-bottom: 0.35rem;
        color: #ffffff;
        letter-spacing: -0.02em;
    }
    .hero-subtitle {
        font-size: 0.95rem;
        color: #94a3b8;
        margin-bottom: 0.8rem;
    }
    .hero-badge {
        display: inline-block;
        background-color: rgba(37, 99, 235, 0.3);
        border: 1px solid rgba(147, 197, 253, 0.4);
        color: #bfdbfe;
        padding: 4px 12px;
        border-radius: 9999px;
        font-size: 0.8rem;
        font-weight: 600;
    }

    /* Metric Cards */
    [data-testid="stMetric"] {
        background: white;
        padding: 1.1rem 1.3rem;
        border-radius: 12px;
        border: 1px solid #e2e8f0;
        box-shadow: 0 2px 6px rgba(0, 0, 0, 0.04);
        transition: transform 0.15s ease, box-shadow 0.15s ease;
    }
    [data-testid="stMetric"]:hover {
        transform: translateY(-2px);
        box-shadow: 0 6px 16px rgba(0, 0, 0, 0.08);
    }
    [data-testid="stMetricLabel"] {
        color: #64748b;
        font-size: 0.85rem;
        font-weight: 600;
    }
    [data-testid="stMetricValue"] {
        color: #0f172a;
        font-weight: 800;
    }

    /* Highlight Card */
    .highlight-card {
        background: #eff6ff;
        border: 1px solid #bfdbfe;
        border-radius: 12px;
        padding: 1rem 1.2rem;
        margin-bottom: 1rem;
    }
    
    /* Table Styling */
    .dataframe {
        font-size: 0.9rem !important;
    }
</style>
""", unsafe_allow_html=True)

# ==========================================
# KAMUS & FUNGSI BANTU
# ==========================================
HARI_INDONESIA = {
    0: "Senin",
    1: "Selasa",
    2: "Rabu",
    3: "Kamis",
    4: "Jumat",
    5: "Sabtu",
    6: "Minggu"
}

def format_rupiah(nominal: float) -> str:
    """Format angka menjadi string mata uang Rupiah standar."""
    if pd.isna(nominal):
        return "Rp 0"
    return f"Rp {nominal:,.0f}".replace(",", ".")

def format_rupiah_sen(nominal: float) -> str:
    """Format angka dengan 2 desimal jika diperlukan."""
    if pd.isna(nominal):
        return "Rp 0,00"
    # Format ribuan dengan koma, lalu ubah ke standar Indonesia
    formatted = f"{nominal:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"Rp {formatted}"

def get_nama_hari(tgl) -> str:
    """Mengembalikan nama hari dalam Bahasa Indonesia dari objek tanggal/string."""
    if tgl is None or pd.isna(tgl):
        return ""
    if isinstance(tgl, str):
        try:
            tgl = datetime.date.fromisoformat(tgl)
        except Exception:
            try:
                tgl = pd.to_datetime(tgl).date()
            except Exception:
                return ""
    if isinstance(tgl, (datetime.date, datetime.datetime)):
        return HARI_INDONESIA.get(tgl.weekday(), "")
    return ""

def parse_time_to_hours(val) -> float:
    """
    Mengonversi string waktu (HH:MM / H:MM / HH.MM), datetime.time,
    menjadi float desimal jam sejak 00:00 (misal: '08:30' -> 8.5).
    """
    if val is None or pd.isna(val):
        raise ValueError("Jam tidak boleh kosong")
    
    if isinstance(val, (datetime.time, datetime.datetime)):
        return val.hour + (val.minute / 60.0) + (val.second / 3600.0)
    
    val_str = str(val).strip()
    if not val_str or val_str.lower() in ["nan", "none", "nat", ""]:
        raise ValueError("Jam tidak boleh kosong")
    
    val_str = val_str.replace(".", ":")
    parts = val_str.split(":")
    if len(parts) >= 2:
        try:
            h = int(parts[0])
            m = int(parts[1])
            s = int(parts[2]) if len(parts) > 2 else 0
            if not (0 <= h <= 24 and 0 <= m < 60 and 0 <= s < 60):
                raise ValueError(f"Waktu {h:02d}:{m:02d} di luar batas valid (00:00 - 23:59)")
            return h + (m / 60.0) + (s / 3600.0)
        except ValueError as err:
            raise ValueError(f"Format jam '{val}' tidak valid. Gunakan format HH:MM (contoh: 08:00)") from err
            
    raise ValueError(f"Format jam '{val}' tidak valid. Gunakan format HH:MM (contoh: 08:00)")

# ==========================================
# LOGIKA BISNIS LEMBUR (PP NO. 35 TAHUN 2021)
# ==========================================
def hitung_lembur_hari_kerja(t_masuk: float, t_keluar: float, std_masuk: float = 8.0, std_keluar: float = 17.0):
    """
    Aturan Hari Kerja (Senin - Jumat) Sistem 5 Hari Kerja:
    - Jam Kerja Standar: 08:00 - 17:00 (8 jam kerja + 1 jam istirahat).
    - Waktu di luar 08:00 s/d 17:00 dihitung sebagai lembur.
    - Faktor Pengali:
      * Jam ke-1: 1.5x upah per jam
      * Jam ke-2 dan seterusnya: 2.0x upah per jam
    """
    # Penanganan jika shift lewat tengah malam (cross-midnight)
    if t_keluar < t_masuk:
        t_keluar_calc = t_keluar + 24.0
    else:
        t_keluar_calc = t_keluar

    durasi_aktual = t_keluar_calc - t_masuk

    # Lembur sebelum jam masuk standar (misal masuk jam 06:00 -> lembur 2 jam)
    lembur_pagi = max(0.0, min(std_masuk, t_keluar_calc) - t_masuk)
    # Lembur setelah jam pulang standar (misal pulang jam 20:00 -> lembur 3 jam)
    lembur_sore = max(0.0, t_keluar_calc - max(std_keluar, t_masuk))
    total_lembur = lembur_pagi + lembur_sore

    if total_lembur <= 0:
        return durasi_aktual, 0.0, 0.0, "-", "Jam kerja standar (tanpa lembur)"

    # Hitung faktor pengali bertingkat
    if total_lembur <= 1.0:
        pengali = total_lembur * 1.5
        rincian = f"{total_lembur:.2f}j × 1.5x"
    else:
        jam_1 = 1.0
        jam_sisa = total_lembur - 1.0
        pengali = (jam_1 * 1.5) + (jam_sisa * 2.0)
        rincian = f"1j × 1.5x + {jam_sisa:.2f}j × 2.0x"

    catatan = []
    if lembur_pagi > 0:
        catatan.append(f"Lembur pagi {lembur_pagi:.2f}j")
    if lembur_sore > 0:
        catatan.append(f"Lembur sore {lembur_sore:.2f}j")
    info = ", ".join(catatan)

    return durasi_aktual, total_lembur, pengali, rincian, info

def hitung_lembur_hari_libur(t_masuk: float, t_keluar: float):
    """
    Aturan Hari Libur (Sabtu, Minggu, Hari Libur Resmi) Sistem 5 Hari Kerja:
    - Total jam kerja aktual langsung dipotong 1 jam untuk istirahat.
    - Faktor Pengali (PP 35/2021 Pasal 31 ayat 1 huruf b):
      * 8 jam pertama: 2.0x upah per jam
      * Jam ke-9: 3.0x upah per jam
      * Jam ke-10 dan seterusnya: 4.0x upah per jam
    """
    if t_keluar < t_masuk:
        t_keluar_calc = t_keluar + 24.0
    else:
        t_keluar_calc = t_keluar

    durasi_aktual = t_keluar_calc - t_masuk
    # Dipotong 1 jam istirahat sesuai instruksi
    jam_lembur = max(0.0, durasi_aktual - 1.0)

    if jam_lembur <= 0:
        return durasi_aktual, 0.0, 0.0, "-", "Durasi ≤ 1 jam (terpotong istirahat 1 jam)"

    # Pengali bertingkat hari libur 5 hari kerja
    if jam_lembur <= 8.0:
        pengali = jam_lembur * 2.0
        rincian = f"{jam_lembur:.2f}j × 2.0x"
    elif jam_lembur <= 9.0:
        jam_8 = 8.0
        jam_9 = jam_lembur - 8.0
        pengali = (8.0 * 2.0) + (jam_9 * 3.0)
        rincian = f"8j × 2.0x + {jam_9:.2f}j × 3.0x"
    else:
        jam_8 = 8.0
        jam_9 = 1.0
        jam_10_plus = jam_lembur - 9.0
        pengali = (8.0 * 2.0) + (1.0 * 3.0) + (jam_10_plus * 4.0)
        rincian = f"8j × 2.0x + 1j × 3.0x + {jam_10_plus:.2f}j × 4.0x"

    info = f"Aktual {durasi_aktual:.2f}j - 1j istirahat = {jam_lembur:.2f}j lembur"
    return durasi_aktual, jam_lembur, pengali, rincian, info

# ==========================================
# GENERATOR LAPORAN EXCEL (.XLSX) PROFESIONAL
# ==========================================
def generate_excel_report(df_hasil: pd.DataFrame, summary_data: dict, employee_info: dict) -> io.BytesIO:
    """
    Membuat file Excel (.xlsx) dengan tata letak profesional,
    palet warna elegan, border, dan formula SUM Excel otomatis.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap Lembur Karyawan"
    ws.views.sheetView[0].showGridLines = True

    # Palet Warna Elegan (Royal Navy / Slate)
    NAVY_DARK = "0F172A"       # Header tabel utama
    NAVY_TITLE = "1E3A8A"      # Judul Laporan
    ZEBRA_FILL = "F8FAFC"      # Baris bergantian
    TOTAL_BG = "E2E8F0"        # Baris Total
    METADATA_BG = "F1F5F9"     # Background informasi
    BORDER_COLOR = "CBD5E1"    # Garis tepi sel

    font_title = Font(name="Segoe UI", size=16, bold=True, color=NAVY_TITLE)
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    font_bold = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    font_regular = Font(name="Segoe UI", size=10, color="1E293B")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_total = Font(name="Segoe UI", size=11, bold=True, color="0F172A")

    fill_header = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
    fill_zebra = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
    fill_total = PatternFill(start_color=TOTAL_BG, end_color=TOTAL_BG, fill_type="solid")
    fill_meta = PatternFill(start_color=METADATA_BG, end_color=METADATA_BG, fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="thin", color=BORDER_COLOR),
        bottom=Side(style="thin", color=BORDER_COLOR)
    )
    border_total = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="medium", color="0F172A"),
        bottom=Side(style="double", color="0F172A")
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # 1. Header Judul
    ws["B1"] = "REKAPITULASI GAJI DAN UPAH LEMBUR KARYAWAN"
    ws["B1"].font = font_title
    ws["B2"] = "Perhitungan Sesuai Aturan Depnaker (PP No. 35 Tahun 2021 - Sistem 5 Hari Kerja)"
    ws["B2"].font = font_subtitle

    # 2. Kotak Metadata Informasi
    meta_box = [
        ("B4", "Nama Karyawan:", "C4", employee_info.get("nama", "-")),
        ("B5", "Jabatan / Divisi:", "C5", employee_info.get("jabatan", "-")),
        ("B6", "Periode Penggajian:", "C6", employee_info.get("periode", "-")),
        ("F4", "Gaji Pokok / UMK:", "G4", summary_data.get("gaji_pokok", 0)),
        ("F5", "Upah Per Jam (1/173):", "G5", summary_data.get("upah_per_jam", 0)),
        ("F6", "Tanggal Cetak:", "G6", datetime.datetime.now().strftime("%d/%m/%Y %H:%M"))
    ]

    for lbl_cell, label, val_cell, val in meta_box:
        ws[lbl_cell] = label
        ws[lbl_cell].font = font_bold
        ws[val_cell] = val
        if "Gaji" in label:
            ws[val_cell].font = font_bold
            ws[val_cell].number_format = '"Rp"#,##0'
        elif "Upah Per Jam" in label:
            ws[val_cell].font = font_regular
            ws[val_cell].number_format = '"Rp"#,##0.00'
        else:
            ws[val_cell].font = font_regular

    # 3. Header Tabel (Mulai Baris 9, Kolom B)
    start_row = 9
    col_defs = [
        ("No", 6, align_center),
        ("Tanggal", 14, align_center),
        ("Hari", 12, align_center),
        ("Keterangan", 22, align_left),
        ("Jam Masuk", 13, align_center),
        ("Jam Keluar", 13, align_center),
        ("Jam Aktual", 14, align_right),
        ("Jam Lembur", 14, align_right),
        ("Pengali (x)", 13, align_right),
        ("Rincian Pengali", 32, align_left),
        ("Nominal Lembur (Rp)", 22, align_right),
        ("Catatan / Keterangan", 30, align_left)
    ]

    ws.column_dimensions["A"].width = 3

    for idx, (col_name, col_width, col_align) in enumerate(col_defs, start=2):
        col_letter = get_column_letter(idx)
        cell = ws.cell(row=start_row, column=idx, value=col_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
        ws.column_dimensions[col_letter].width = col_width

    # 4. Pengisian Baris Data
    current_row = start_row + 1
    for i, row in df_hasil.iterrows():
        is_zebra = (i % 2 == 1)
        row_fill = fill_zebra if is_zebra else None

        row_values = [
            (i + 1, align_center, None),
            (str(row["Tanggal"]), align_center, None),
            (str(row["Hari"]), align_center, None),
            (str(row["Keterangan"]), align_left, None),
            (str(row["Jam Masuk"]), align_center, None),
            (str(row["Jam Keluar"]), align_center, None),
            (float(row["Jam Aktual"]), align_right, "0.00"),
            (float(row["Jam Lembur"]), align_right, "0.00"),
            (float(row["Faktor Pengali"]), align_right, "0.00"),
            (str(row["Rincian Pengali"]), align_left, None),
            (float(row["Nominal Lembur"]), align_right, '"Rp"#,##0'),
            (str(row.get("Catatan", "")), align_left, None)
        ]

        for col_idx, (val, alignment, num_fmt) in enumerate(row_values, start=2):
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.font = font_regular
            c.alignment = alignment
            c.border = thin_border
            if row_fill:
                c.fill = row_fill
            if num_fmt:
                c.number_format = num_fmt

        current_row += 1

    # 5. Baris Total Keseluruhan
    total_row = current_row
    ws.merge_cells(start_row=total_row, start_column=2, end_row=total_row, end_column=7)
    lbl_tot = ws.cell(row=total_row, column=2, value="TOTAL KESELURUHAN")
    lbl_tot.font = font_total
    lbl_tot.alignment = Alignment(horizontal="right", vertical="center")

    for col in range(2, 8):
        c = ws.cell(row=total_row, column=col)
        c.fill = fill_total
        c.border = border_total

    # SUM Jam Aktual (Kolom H / 8)
    col_h = get_column_letter(8)
    c_sum_akt = ws.cell(row=total_row, column=8)
    c_sum_akt.value = f"=SUM({col_h}{start_row + 1}:{col_h}{total_row - 1})"
    c_sum_akt.font = font_total
    c_sum_akt.alignment = align_right
    c_sum_akt.number_format = "0.00"
    c_sum_akt.fill = fill_total
    c_sum_akt.border = border_total

    # SUM Jam Lembur (Kolom I / 9)
    col_i = get_column_letter(9)
    c_sum_lbr = ws.cell(row=total_row, column=9)
    c_sum_lbr.value = f"=SUM({col_i}{start_row + 1}:{col_i}{total_row - 1})"
    c_sum_lbr.font = font_total
    c_sum_lbr.alignment = align_right
    c_sum_lbr.number_format = "0.00"
    c_sum_lbr.fill = fill_total
    c_sum_lbr.border = border_total

    # SUM Pengali (Kolom J / 10)
    col_j = get_column_letter(10)
    c_sum_png = ws.cell(row=total_row, column=10)
    c_sum_png.value = f"=SUM({col_j}{start_row + 1}:{col_j}{total_row - 1})"
    c_sum_png.font = font_total
    c_sum_png.alignment = align_right
    c_sum_png.number_format = "0.00"
    c_sum_png.fill = fill_total
    c_sum_png.border = border_total

    # Kosong Rincian (Kolom K / 11)
    c_blank = ws.cell(row=total_row, column=11, value="")
    c_blank.fill = fill_total
    c_blank.border = border_total

    # SUM Nominal Lembur (Kolom L / 12)
    col_l = get_column_letter(12)
    c_sum_nom = ws.cell(row=total_row, column=12)
    c_sum_nom.value = f"=SUM({col_l}{start_row + 1}:{col_l}{total_row - 1})"
    c_sum_nom.font = font_total
    c_sum_nom.alignment = align_right
    c_sum_nom.number_format = '"Rp"#,##0'
    c_sum_nom.fill = fill_total
    c_sum_nom.border = border_total

    # Kosong Catatan (Kolom M / 13)
    c_note = ws.cell(row=total_row, column=13, value="")
    c_note.fill = fill_total
    c_note.border = border_total

    # 6. Ringkasan Take Home Pay di Bawah Tabel
    summary_start = total_row + 2
    ws.cell(row=summary_start, column=10, value="Gaji Pokok / UMK:").font = font_bold
    ws.cell(row=summary_start, column=12, value=summary_data.get("gaji_pokok", 0)).font = font_bold
    ws.cell(row=summary_start, column=12).number_format = '"Rp"#,##0'
    ws.cell(row=summary_start, column=12).alignment = align_right

    ws.cell(row=summary_start + 1, column=10, value="Total Uang Lembur:").font = font_bold
    ws.cell(row=summary_start + 1, column=12, value=summary_data.get("total_uang_lembur", 0)).font = font_bold
    ws.cell(row=summary_start + 1, column=12).number_format = '"Rp"#,##0'
    ws.cell(row=summary_start + 1, column=12).alignment = align_right

    ws.cell(row=summary_start + 2, column=10, value="Take Home Pay (Kotor):").font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")
    c_thp = ws.cell(row=summary_start + 2, column=12, value=summary_data.get("take_home_pay", 0))
    c_thp.font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")
    c_thp.number_format = '"Rp"#,##0'
    c_thp.alignment = align_right

    # 7. Catatan Kaki Regulasi
    foot_row = summary_start + 4
    ws.cell(row=foot_row, column=2, value="Dasar Hukum & Ketentuan Lembur (PP No. 35 Tahun 2021):").font = font_bold
    ws.cell(row=foot_row + 1, column=2, value="1. Sistem 5 Hari Kerja: Jam kerja normal 08:00 - 17:00 (Senin - Jumat). Lembur dihitung di luar rentang jam tersebut.").font = font_subtitle
    ws.cell(row=foot_row + 2, column=2, value="2. Hari Kerja: Jam ke-1 dibayar 1.5x upah per jam; Jam ke-2 dan seterusnya dibayar 2.0x upah per jam.").font = font_subtitle
    ws.cell(row=foot_row + 3, column=2, value="3. Hari Libur: Jam aktual dipotong istirahat 1 jam. 8 jam pertama 2.0x; Jam ke-9 3.0x; Jam ke-10 dst 4.0x.").font = font_subtitle
    ws.cell(row=foot_row + 4, column=2, value="4. Upah per jam = 1/173 x Gaji Pokok sebulan (Pasal 32 ayat 2).").font = font_subtitle

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer

# ==========================================
# DATA AWAL & HELPER STATE
# ==========================================
def get_sample_dataframe():
    """Mengembalikan data contoh 1 minggu yang realistis."""
    today = datetime.date.today()
    # Cari hari Senin terdekat di minggu berjalan
    senin = today - datetime.timedelta(days=today.weekday())
    sample_rows = [
        {"Tanggal": senin + datetime.timedelta(days=0), "Hari": "Senin", "Keterangan": "Hari Kerja", "Jam Masuk": "08:00", "Jam Keluar": "17:00", "Catatan": "Jam kerja normal"},
        {"Tanggal": senin + datetime.timedelta(days=1), "Hari": "Selasa", "Keterangan": "Hari Kerja", "Jam Masuk": "08:00", "Jam Keluar": "20:00", "Catatan": "Lembur sore 3 jam"},
        {"Tanggal": senin + datetime.timedelta(days=2), "Hari": "Rabu", "Keterangan": "Hari Kerja", "Jam Masuk": "06:00", "Jam Keluar": "17:00", "Catatan": "Lembur pagi 2 jam"},
        {"Tanggal": senin + datetime.timedelta(days=3), "Hari": "Kamis", "Keterangan": "Hari Kerja", "Jam Masuk": "07:00", "Jam Keluar": "19:30", "Catatan": "Lembur pagi 1j + sore 2.5j"},
        {"Tanggal": senin + datetime.timedelta(days=4), "Hari": "Jumat", "Keterangan": "Hari Kerja", "Jam Masuk": "08:00", "Jam Keluar": "17:00", "Catatan": "Jam kerja normal"},
        {"Tanggal": senin + datetime.timedelta(days=5), "Hari": "Sabtu", "Keterangan": "Hari Libur / Weekend", "Jam Masuk": "08:00", "Jam Keluar": "17:00", "Catatan": "Lembur weekend (potong 1j istirahat)"},
        {"Tanggal": senin + datetime.timedelta(days=6), "Hari": "Minggu", "Keterangan": "Hari Libur / Weekend", "Jam Masuk": "09:00", "Jam Keluar": "14:00", "Catatan": "Lembur weekend 5j (efektif 4j)"}
    ]
    return pd.DataFrame(sample_rows)

def get_blank_dataframe():
    """Mengembalikan 1 baris kosong sebagai permulaan."""
    today = datetime.date.today()
    day_name = get_nama_hari(today)
    is_weekend = today.weekday() in [5, 6]
    return pd.DataFrame([{
        "Tanggal": today,
        "Hari": day_name,
        "Keterangan": "Hari Libur / Weekend" if is_weekend else "Hari Kerja",
        "Jam Masuk": "08:00",
        "Jam Keluar": "17:00",
        "Catatan": ""
    }])

def get_monthly_template(year: int, month: int):
    """Menghasilkan baris tanggal untuk 1 bulan penuh otomatis."""
    num_days = calendar.monthrange(year, month)[1]
    rows = []
    for day in range(1, num_days + 1):
        dt = datetime.date(year, month, day)
        day_name = get_nama_hari(dt)
        is_weekend = dt.weekday() in [5, 6]
        rows.append({
            "Tanggal": dt,
            "Hari": day_name,
            "Keterangan": "Hari Libur / Weekend" if is_weekend else "Hari Kerja",
            "Jam Masuk": "08:00" if not is_weekend else "",
            "Jam Keluar": "17:00" if not is_weekend else "",
            "Catatan": ""
        })
    return pd.DataFrame(rows)

# Inisialisasi Session State
if "df_input" not in st.session_state:
    st.session_state.df_input = get_sample_dataframe()

if "calc_results" not in st.session_state:
    st.session_state.calc_results = None

# ==========================================
# SIDEBAR: PARAMETER UTAMA & PENGATURAN
# ==========================================
with st.sidebar:
    st.image("https://cdn-icons-png.flaticon.com/512/3135/3135706.png", width=64)
    st.title("⚙️ Parameter Lembur")
    st.caption("PP No. 35 Tahun 2021 (Sistem 5 Hari Kerja)")
    st.divider()

    st.subheader("💰 Pengaturan Gaji")
    # Default 5.247.870 sesuai spesifikasi
    gaji_pokok = st.number_input(
        "Gaji Pokok / UMK (Rp)",
        min_value=1_000_000.0,
        max_value=200_000_000.0,
        value=5_247_870.0,
        step=50_000.0,
        format="%.0f",
        help="Gaji Pokok atau UMK bulanan yang menjadi dasar perhitungan lembur."
    )
    st.markdown(f"**Terbilang/Format:** `{format_rupiah(gaji_pokok)}`")

    # Rumus Upah Per Jam = Gaji Pokok / 173 (PP 35/2021 Pasal 32 ayat 2)
    upah_per_jam = gaji_pokok / 173.0

    st.markdown(f"""
    <div class="highlight-card">
        <div style="font-size:0.8rem; color:#475569; font-weight:600;">UPAH PER JAM (1/173)</div>
        <div style="font-size:1.35rem; color:#1e3a8a; font-weight:800;">{format_rupiah_sen(upah_per_jam)}</div>
        <div style="font-size:0.75rem; color:#64748b; margin-top:4px;">Dasar: Pasal 32 ayat (2) PP No. 35/2021</div>
    </div>
    """, unsafe_allow_html=True)

    st.divider()
    st.subheader("👤 Data Karyawan (Untuk Excel)")
    nama_karyawan = st.text_input("Nama Karyawan", value="Budi Santoso")
    jabatan_karyawan = st.text_input("Jabatan / Departemen", value="Staff Operasional")
    periode_laporan = st.text_input("Periode Laporan", value=datetime.date.today().strftime("%B %Y"))

    st.divider()
    st.subheader("⚡ Aksi Cepat Preset Data")
    col_pre1, col_pre2 = st.columns(2)
    with col_pre1:
        if st.button("✨ Contoh Data", use_container_width=True, help="Muat data contoh 1 minggu kerja & lembur"):
            st.session_state.df_input = get_sample_dataframe()
            st.session_state.calc_results = None
            st.rerun()

    with col_pre2:
        if st.button("🗑️ Kosongkan", use_container_width=True, help="Hapus seluruh data menjadi tabel kosong"):
            st.session_state.df_input = get_blank_dataframe()
            st.session_state.calc_results = None
            st.rerun()

    # Generator 1 Bulan Penuh
    with st.expander("📅 Buat Template 1 Bulan"):
        tgl_skrg = datetime.date.today()
        pilih_bulan = st.selectbox("Pilih Bulan", list(range(1, 13)), index=tgl_skrg.month - 1)
        pilih_tahun = st.selectbox("Pilih Tahun", [tgl_skrg.year - 1, tgl_skrg.year, tgl_skrg.year + 1], index=1)
        if st.button("Generate Template Bulan", use_container_width=True):
            st.session_state.df_input = get_monthly_template(pilih_tahun, pilih_bulan)
            st.session_state.calc_results = None
            st.success(f"Berhasil membuat template untuk bulan {calendar.month_name[pilih_bulan]} {pilih_tahun}!")
            st.rerun()

# ==========================================
# MAIN CONTENT AREA
# ==========================================

# 1. Hero Header
st.markdown("""
<div class="hero-container">
    <div class="hero-badge">⚖️ KEMNAKER RI • PP NO. 35 TAHUN 2021</div>
    <div class="hero-title">Kalkulator Gaji & Lembur Karyawan</div>
    <div class="hero-subtitle">
        Sistem Perhitungan Lembur Otomatis Khusus 5 Hari Kerja (40 Jam/Minggu) dengan Ekspor Laporan Excel (.xlsx) Resmi
    </div>
</div>
""", unsafe_allow_html=True)

# Tab Navigasi
tab_input, tab_panduan = st.tabs(["📝 Input Data & Kalkulasi", "📖 Panduan Hukum & Formula PP 35/2021"])

with tab_input:
    st.subheader("1. Input Jadwal Kehadiran & Jam Kerja")
    st.caption("Silakan isi atau sesuaikan tabel di bawah ini. Anda dapat menambah, menghapus, atau mengubah baris secara dinamis.")

    # Konfigurasi Kolom st.data_editor
    column_config = {
        "Tanggal": st.column_config.DateColumn(
            "Tanggal",
            help="Pilih tanggal kehadiran",
            format="YYYY-MM-DD",
            required=True
        ),
        "Hari": st.column_config.TextColumn(
            "Hari",
            help="Nama hari dalam Bahasa Indonesia (akan diselaraskan otomatis jika tanggal berubah)",
            disabled=False
        ),
        "Keterangan": st.column_config.SelectboxColumn(
            "Jenis Hari",
            help="Pilih apakah Hari Kerja reguler atau Hari Libur / Weekend / Hari Besar",
            options=["Hari Kerja", "Hari Libur / Weekend"],
            required=True
        ),
        "Jam Masuk": st.column_config.TextColumn(
            "Jam Masuk",
            help="Waktu mulai kerja format HH:MM (contoh: 08:00)",
            required=True
        ),
        "Jam Keluar": st.column_config.TextColumn(
            "Jam Keluar",
            help="Waktu selesai kerja format HH:MM (contoh: 17:00 atau 20:00)",
            required=True
        ),
        "Catatan": st.column_config.TextColumn(
            "Catatan",
            help="Keterangan pekerjaan atau nomor surat tugas lembur (opsional)"
        )
    }

    # Render data_editor
    edited_df = st.data_editor(
        st.session_state.df_input,
        column_config=column_config,
        num_rows="dynamic",
        use_container_width=True,
        key="editor_kehadiran"
    )

    # Sinkronisasi nama hari otomatis jika tanggal diedit
    try:
        updated_hari = []
        for _, r in edited_df.iterrows():
            tgl_val = r.get("Tanggal")
            h_val = get_nama_hari(tgl_val)
            updated_hari.append(h_val if h_val else r.get("Hari", ""))
        edited_df["Hari"] = updated_hari
    except Exception:
        pass

    # Simpan kembali ke session_state
    st.session_state.df_input = edited_df

    st.markdown("<br>", unsafe_allow_html=True)

    # Tombol Aksi Utama
    col_btn, col_info_btn = st.columns([1, 3])
    with col_btn:
        tombol_kalkulasi = st.button("🚀 Kalkulasi Lembur", type="primary", use_container_width=True)
    with col_info_btn:
        st.caption("Klik tombol **Kalkulasi Lembur** untuk memproses data sesuai aturan PP 35/2021.")

    # ==========================================
    # LOGIKA PROSES KALKULASI & ERROR HANDLING
    # ==========================================
    if tombol_kalkulasi:
        errors = []
        hasil_list = []

        if edited_df.empty:
            st.error("⚠️ Data editor masih kosong. Silakan tambahkan minimal 1 baris kehadiran.")
        else:
            for idx, row in edited_df.iterrows():
                row_num = idx + 1
                tgl = row.get("Tanggal")
                hari = row.get("Hari")
                keterangan = row.get("Keterangan", "Hari Kerja")
                jam_masuk_raw = row.get("Jam Masuk")
                jam_keluar_raw = row.get("Jam Keluar")
                catatan = row.get("Catatan", "")

                # Validasi jika baris kosong sepenuhnya
                if (tgl is None or pd.isna(tgl)) and not jam_masuk_raw and not jam_keluar_raw:
                    continue

                # Validasi Tanggal
                if tgl is None or pd.isna(tgl):
                    errors.append(f"Baris #{row_num}: Tanggal tidak boleh kosong.")
                    continue

                if not hari:
                    hari = get_nama_hari(tgl)

                # Validasi Jam Masuk & Jam Keluar dengan try-except
                try:
                    t_masuk = parse_time_to_hours(jam_masuk_raw)
                except ValueError as ve:
                    errors.append(f"Baris #{row_num} (Tanggal {tgl}): Jam Masuk error -> {str(ve)}")
                    continue

                try:
                    t_keluar = parse_time_to_hours(jam_keluar_raw)
                except ValueError as ve:
                    errors.append(f"Baris #{row_num} (Tanggal {tgl}): Jam Keluar error -> {str(ve)}")
                    continue

                if t_masuk == t_keluar:
                    # Tidak ada durasi kerja
                    durasi_akt = 0.0
                    jam_lbr = 0.0
                    pengali = 0.0
                    rincian = "-"
                    info_det = "Jam masuk dan keluar sama (0 jam)"
                else:
                    if keterangan == "Hari Kerja":
                        durasi_akt, jam_lbr, pengali, rincian, info_det = hitung_lembur_hari_kerja(t_masuk, t_keluar)
                    else:
                        durasi_akt, jam_lbr, pengali, rincian, info_det = hitung_lembur_hari_libur(t_masuk, t_keluar)

                nominal_lembur = round(pengali * upah_per_jam)

                hasil_list.append({
                    "Tanggal": tgl,
                    "Hari": hari,
                    "Keterangan": keterangan,
                    "Jam Masuk": str(jam_masuk_raw).strip(),
                    "Jam Keluar": str(jam_keluar_raw).strip(),
                    "Jam Aktual": round(durasi_akt, 2),
                    "Jam Lembur": round(jam_lbr, 2),
                    "Faktor Pengali": round(pengali, 2),
                    "Rincian Pengali": rincian,
                    "Nominal Lembur": nominal_lembur,
                    "Catatan": f"{info_det}. {catatan}".strip() if catatan else info_det
                })

            if errors:
                st.error("⚠️ Ditemukan kesalahan pada input data:")
                for err in errors:
                    st.write(f"- {err}")
                st.info("💡 Pastikan format jam menggunakan format HH:MM (contoh: 08:00, 17:30, 20:00).")

            if hasil_list:
                df_hasil = pd.DataFrame(hasil_list)
                total_jam_lembur = df_hasil["Jam Lembur"].sum()
                total_uang_lembur = df_hasil["Nominal Lembur"].sum()
                take_home_pay = gaji_pokok + total_uang_lembur

                st.session_state.calc_results = {
                    "df_hasil": df_hasil,
                    "gaji_pokok": gaji_pokok,
                    "upah_per_jam": upah_per_jam,
                    "total_jam_lembur": total_jam_lembur,
                    "total_uang_lembur": total_uang_lembur,
                    "take_home_pay": take_home_pay
                }

    # ==========================================
    # TAMPILAN DASHBOARD SUMMARY & HASIL
    # ==========================================
    if st.session_state.calc_results is not None:
        res = st.session_state.calc_results
        df_res = res["df_hasil"]

        st.markdown("<br><hr>", unsafe_allow_html=True)
        st.subheader("2. Dashboard Summary & Ringkasan Gaji")

        # 4 Metric Cards Utama
        col_m1, col_m2, col_m3, col_m4 = st.columns(4)
        with col_m1:
            st.metric(
                label="⏱️ Total Jam Lembur",
                value=f"{res['total_jam_lembur']:.2f} Jam",
                help="Akumulasi total jam lembur yang sah sesuai PP 35/2021"
            )
        with col_m2:
            st.metric(
                label="💵 Total Uang Lembur",
                value=format_rupiah(res["total_uang_lembur"]),
                help="Total kompensasi lembur yang berhak diterima karyawan"
            )
        with col_m3:
            st.metric(
                label="💰 Gaji Pokok (UMK)",
                value=format_rupiah(res["gaji_pokok"]),
                help="Gaji Pokok bulanan karyawan"
            )
        with col_m4:
            st.metric(
                label="💳 Take Home Pay (Kotor)",
                value=format_rupiah(res["take_home_pay"]),
                delta=f"+ {format_rupiah(res['total_uang_lembur'])} (Lembur)",
                help="Total penerimaan kotor = Gaji Pokok + Total Uang Lembur"
            )

        st.markdown("<br>", unsafe_allow_html=True)
        st.subheader("3. Rincian Hasil Perhitungan Lembur Per Hari")

        # Tampilkan DataFrame Hasil dengan formatting mata uang
        df_display = df_res.copy()
        df_display["Nominal Lembur (Rp)"] = df_display["Nominal Lembur"].apply(format_rupiah)
        
        cols_order = [
            "Tanggal", "Hari", "Keterangan", "Jam Masuk", "Jam Keluar",
            "Jam Aktual", "Jam Lembur", "Faktor Pengali", "Rincian Pengali",
            "Nominal Lembur (Rp)", "Catatan"
        ]
        
        st.dataframe(
            df_display[cols_order],
            use_container_width=True,
            hide_index=True
        )

        # ==========================================
        # FITUR EKSPOR REKAPAN EXCEL (.XLSX)
        # ==========================================
        st.markdown("<br>", unsafe_allow_html=True)
        st.subheader("4. Ekspor Laporan Rekap ke Excel (.xlsx)")

        emp_info = {
            "nama": nama_karyawan,
            "jabatan": jabatan_karyawan,
            "periode": periode_laporan
        }

        # Generate file Excel dalam memori (io.BytesIO)
        excel_data = generate_excel_report(df_res, res, emp_info)

        file_name_excel = f"Laporan_Lembur_{nama_karyawan.replace(' ', '_')}_{periode_laporan.replace(' ', '_')}.xlsx"

        col_dl1, col_dl2 = st.columns([1, 2])
        with col_dl1:
            st.download_button(
                label="📥 Unduh Rekap Excel (.xlsx)",
                data=excel_data,
                file_name=file_name_excel,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True
            )
        with col_dl2:
            st.caption(
                f"File: **{file_name_excel}** sudah diformat dengan standar akuntansi resmi, header warna Navy, "
                "rumus SUM otomatis, dan format mata uang Rupiah siap cetak/arsip."
            )

        # Visualisasi Ringan (Grafik Jam Lembur)
        with st.expander("📊 Grafik Visualisasi Lembur Harian"):
            col_chart1, col_chart2 = st.columns(2)
            with col_chart1:
                st.caption("Distribusi Jam Lembur Per Tanggal")
                chart_df = df_res.set_index("Tanggal")[["Jam Lembur"]]
                st.bar_chart(chart_df)
            with col_chart2:
                st.caption("Komparasi Uang Lembur Berdasarkan Jenis Hari")
                group_hari = df_res.groupby("Keterangan")["Nominal Lembur"].sum().reset_index()
                st.bar_chart(data=group_hari, x="Keterangan", y="Nominal Lembur")

# ==========================================
# TAB 2: PANDUAN HUKUM PP NO. 35 TAHUN 2021
# ==========================================
with tab_panduan:
    st.subheader("📖 Dasar Hukum & Formula Depnaker (PP No. 35 Tahun 2021)")
    st.markdown("""
    Perhitungan upah kerja lembur diatur secara ketat dalam **Peraturan Pemerintah Republik Indonesia Nomor 35 Tahun 2021**
    tentang Perjanjian Kerja Waktu Tertentu, Alih Daya, Waktu Kerja dan Waktu Istirahat, dan Pemutusan Hubungan Kerja.

    ---
    ### 1. Rumus Upah Per Jam (Pasal 32 Ayat 2)
    > *"Dalam hal Upah sebulan terdiri atas Upah pokok dan tunjangan tetap atau Upah tanpa tunjangan, dasar perhitungan Upah sejam adalah **1/173 (satu per seratus tujuh puluh tiga)** kali Upah sebulan."*

    $$\\text{Upah Sejam} = \\frac{\\text{Gaji Pokok / UMK}}{173}$$

    ---
    ### 2. Aturan Hari Kerja Reguler (Pasal 31 Ayat 1 Huruf a)
    Untuk sistem **5 (lima) hari kerja** dan 40 jam seminggu, jam kerja normal adalah pukul **08:00 - 17:00** (8 jam kerja + 1 jam istirahat).
    Pekerjaan di luar jam kerja reguler tersebut dikategorikan sebagai kerja lembur dengan tarif faktor pengali:
    - **Jam lembur pertama (ke-1)**: Dibayar sebesar **1,5 kali (1.5x)** Upah sejam.
    - **Jam lembur berikutnya (ke-2 dan seterusnya)**: Dibayar sebesar **2 kali (2.0x)** Upah sejam.

    *Contoh Kasus:*
    - Karyawan masuk pukul **06:00** dan pulang pukul **20:00**.
    - Lembur pagi (06:00 - 08:00) = 2 jam.
    - Lembur sore (17:00 - 20:00) = 3 jam.
    - Total Lembur = **5 jam**.
    - Perhitungan Pengali:
      * Jam ke-1: $1 \\times 1,5 = 1,5$
      * Jam ke-2 s/d ke-5 (4 jam): $4 \\times 2,0 = 8,0$
      * Total Faktor Pengali = $1,5 + 8,0 = \\mathbf{9,5\\times}$ Upah Sejam.

    ---
    ### 3. Aturan Hari Libur / Istirahat Mingguan (Pasal 31 Ayat 1 Huruf b)
    Untuk perusahaan dengan sistem **5 (lima) hari kerja**, lembur pada hari istirahat mingguan (Sabtu/Minggu) atau hari libur resmi dihitung dengan ketentuan:
    - Jam kerja aktual dipotong **1 jam untuk waktu istirahat**.
    - **8 jam pertama**: Dibayar sebesar **2 kali (2.0x)** Upah sejam.
    - **Jam ke-9**: Dibayar sebesar **3 kali (3.0x)** Upah sejam.
    - **Jam ke-10, ke-11, ke-12**: Dibayar sebesar **4 kali (4.0x)** Upah sejam.

    *Contoh Kasus:*
    - Karyawan bekerja pada hari Sabtu dari pukul **08:00** s/d **18:00** (10 jam aktual).
    - Dipotong 1 jam istirahat $\\rightarrow$ Jam lembur sah = **9 jam**.
    - Perhitungan Pengali:
      * 8 jam pertama: $8 \\times 2,0 = 16,0$
      * Jam ke-9 (1 jam): $1 \\times 3,0 = 3,0$
      * Total Faktor Pengali = $16,0 + 3,0 = \\mathbf{19,0\\times}$ Upah Sejam.
    """)

# Footer Aplikasi
st.markdown("<hr>", unsafe_allow_html=True)
st.markdown(
    "<div style='text-align: center; color: #94a3b8; font-size: 0.85rem;'>"
    "Kalkulator Gaji & Lembur Karyawan • Sesuai PP No. 35 Tahun 2021 • Siap Dideploy di Streamlit Cloud"
    "</div>",
    unsafe_allow_html=True
)
